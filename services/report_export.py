"""Teacher-readable Excel report export without model reasoning."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from services.dashboard_service import DashboardSummary, StudentPerformance


def build_class_report(records: list[StudentPerformance], summary: DashboardSummary) -> bytes:
    workbook = Workbook()
    student_sheet = workbook.active
    student_sheet.title = "Student Summary"
    student_sheet.append(["Student Name", "Class", "Parent/Guardian", "Parent Email", "Assessment", "Marks", "Maximum", "Percentage", "Review Status"])
    for record in records:
        student_sheet.append([record.student_name, record.class_name or "Not available", record.parent_guardian_name or "Not available", record.parent_email or "Not available", record.assessment_name, record.marks, record.maximum, record.percentage, record.review_status])
    question_sheet = workbook.create_sheet("Question-Topic Performance")
    question_sheet.append(["Student", "Question/Topic", "Score", "Maximum", "Percentage", "Status"])
    for record in records:
        for question in record.questions:
            question_sheet.append([record.student_name, question.question_id, question.score, question.maximum, question.percentage, question.status])
    summary_sheet = workbook.create_sheet("Class Summary")
    summary_sheet.append(["Class", "Assessment", "Average", "Highest", "Lowest", "Number processed", "Number requiring review"])
    classes = ", ".join(sorted({record.class_name for record in records if record.class_name})) or "Not available"
    assessments = ", ".join(sorted({record.assessment_name for record in records})) or "Not available"
    summary_sheet.append([classes, assessments, summary.average_score, summary.highest_score, summary.lowest_score, summary.assessments_processed, summary.needs_review])
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(48, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
