from io import BytesIO

from openpyxl import load_workbook

from core.models import VisionExtraction
from core.partial_credit import PartialCreditSuggestion, SuggestionStatus
from core.teacher_review import ReviewAction
from database.audit_store import AuditStore, LifecycleState
from database.db import initialize_database
from services.dashboard_service import dashboard_summary, verified_student_performance
import services.email_service as email_service
from services.email_service import report_email_preview, smtp_ready, student_report_body
from services.report_export import build_class_report


def _verified_record(tmp_path, name="Student", score=8, maximum=10):
    connection = initialize_database(tmp_path / "dashboard.db")
    store = AuditStore(connection)
    extraction = VisionExtraction.model_validate({
        "student": {"name": name, "class_name": "9", "subject": "English"},
        "worksheet_reported_score": {"obtained": score, "maximum": maximum},
        "sections": [{"questions": [{"identifier": "Q1", "student_answer": {"visible_text": "Answer"}, "teacher_marking": {"visible_individual_score": {"obtained": score}}}]}],
    })
    audit_id = store.create_extraction(extraction)
    store.record_suggestion(audit_id, PartialCreditSuggestion(question_identifier="Q1", teacher_awarded_mark=score, rubric_maximum_marks=maximum, evaluation_maximum_marks=maximum, suggested_mark=score, raw_semantic_coverage=score, rationale="Visible evidence.", status=SuggestionStatus.AGREES_WITH_TEACHER))
    store.transition(audit_id, LifecycleState.TEACHER_CONFIRMED, teacher_action=ReviewAction.CONFIRM_EXTRACTED)
    store.transition(audit_id, LifecycleState.VERIFIED, teacher_action=ReviewAction.CONFIRM_EXTRACTED)
    return connection


def test_dashboard_uses_verified_records_only_and_keeps_missing_parent_data_unavailable(tmp_path):
    connection = _verified_record(tmp_path)
    connection.execute("INSERT INTO assessment_audits(extraction_fingerprint, original_extraction_json, lifecycle_state) VALUES ('draft', '{\"sections\":[]}', 'REVIEW_REQUIRED')")
    connection.commit()
    records = verified_student_performance(connection)

    assert len(records) == 1
    assert records[0].student_name == "Student"
    assert records[0].parent_guardian_name is None
    assert records[0].parent_email is None
    assert records[0].percentage == 80
    assert records[0].questions[0].percentage == 80


def test_dashboard_summary_and_excel_report(tmp_path):
    connection = _verified_record(tmp_path, name="One", score=8)
    records = verified_student_performance(connection)
    summary = dashboard_summary(records)
    report = build_class_report(records, summary)
    workbook = load_workbook(BytesIO(report))

    assert summary.average_score == 8
    assert workbook.sheetnames == ["Student Summary", "Question-Topic Performance", "Class Summary"]
    assert workbook["Student Summary"][2][0].value == "One"
    assert workbook["Question-Topic Performance"][2][1].value == "Q1"


def test_email_preview_never_sends_and_smtp_requires_configuration():
    from config.settings import Settings

    preview = report_email_preview("GradeAssist class report", "Teacher summary", ["parent@example.com"])
    assert preview["To"] == "parent@example.com"
    assert "Teacher summary" in preview.get_content()
    assert smtp_ready(Settings(_env_file=None)) is False


def test_persisted_contact_is_used_and_missing_contact_is_never_fabricated(tmp_path):
    connection = _verified_record(tmp_path, name="Sachin")
    connection.execute("INSERT INTO students(name, class_name, parent_guardian_name, parent_email) VALUES (?, ?, ?, ?)", ("Sachin", "9", "Guardian", "guardian@example.com"))
    connection.commit()
    record = verified_student_performance(connection)[0]

    assert record.parent_guardian_name == "Guardian"
    assert record.parent_email == "guardian@example.com"
    body = student_report_body(student_name=record.student_name, assessment_name=record.assessment_name, marks=record.marks, maximum=record.maximum, percentage=record.percentage, verified_status=record.review_status, timestamp=record.assessment_date, strengths=[], needs_attention=[])
    assert "guardian@example.com" not in body
    assert "Strengths supported" not in body


def test_explicit_send_calls_smtp_while_preview_does_not(monkeypatch):
    from config.settings import Settings

    calls = []
    class FakeSMTP:
        def __init__(self, host, port, timeout): calls.append(("connect", host, port))
        def __enter__(self): return self
        def __exit__(self, *args): return None
        def starttls(self): calls.append(("tls",))
        def login(self, username, password): calls.append(("login", username))
        def send_message(self, message): calls.append(("send", message["To"]))

    monkeypatch.setattr(email_service.smtplib, "SMTP", FakeSMTP)
    settings = Settings(_env_file=None, smtp_host="smtp.example.com", smtp_from_email="teacher@example.com")
    report_email_preview("Preview", "Body", ["parent@example.com"])
    assert calls == []
    email_service.send_report_email(settings, ["parent@example.com"], "Report", "Body", b"report")
    assert ("send", "parent@example.com") in calls
